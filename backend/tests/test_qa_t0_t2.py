# SPDX-License-Identifier: GPL-3.0-or-later
"""QA（严过关）对抗性复核测试 —— T0 / T1 / T2 独立验证。

本文件**不由工程师编写**，目的是**尝试证伪**工程师的自测结论，而非复跑其用例。
每条用例上方注明「意图」。命名以 ``test_qa_`` 前缀区分。

覆盖：
- A 头号红线：无效输出**绝不**补默认答案 / 不做大小写空白容错 / 不递归剥围栏。
- A3 ``ValidatedAnswer.score`` 由 value 派生，模型传入的 score 不得生效。
- C 数据导入「绝不静默丢行」+ **真·20,000 行进库并直接用 SQL 复核**。
- G 契约完整性（六个强制模型命名、5 档映射、单题结构性约束、无自动年龄过滤）。
- D 测试库隔离的运行时断言。

验收命令：``cd backend && uv run pytest tests/test_qa_t0_t2.py -q``
"""

from __future__ import annotations

import csv
import io
from decimal import Decimal

import pytest
from pydantic import ValidationError
from sqlalchemy import text

from app.contracts import (
    OPTION_SCORES,
    OPTION_VALUES,
    PURCHASE_INTENT_QUESTION_ID,
    ModelRequest,
    ModelResponse,
    PersonaSnapshot,
    ProductInput,
    QuestionInput,
    QuestionOption,
    RunCreate,
    SurveyInput,
    ValidatedAnswer,
    default_options,
)
from app.inference.validation import InvalidOutputError, parse_and_validate
from app.personas.source import (
    InvalidUserTableError,
    build_snapshots,
    create_import,
    parse_user_table,
)

CSV_MAPPING = {
    "persona_id": "user_id",
    "profile_text_columns": ["age", "city", "gender", "note"],
    "age": "age",
    "city_tier": "city",
}
HEADER = ["user_id", "age", "city", "gender", "note"]


def _csv_bytes(header, data_rows, *, bom: bool = False) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(header)
    for row in data_rows:
        writer.writerow(row)
    return buffer.getvalue().encode("utf-8-sig" if bom else "utf-8")


# ===========================================================================
# A. 头号红线：无效输出必须抛 InvalidOutputError，绝不返回任何答案
# ===========================================================================


# 每条 (raw_text, 描述)。意图：任何一条若**未能抛错**，即证明存在补值/容错路径。
_ILLEGAL_OUTPUTS: list[tuple[object, str]] = [
    # value 类型不对
    ('{"question_id":"purchase_intent","value":3}', "数值型 value=3"),
    ('{"question_id":"purchase_intent","value":true}', "布尔型 value=true"),
    ('{"question_id":"purchase_intent","value":null}', "value=null"),
    ('{"question_id":"purchase_intent","value":[]}', "数组型 value"),
    # value 内容不对（含大小写/空白不得容错）
    ('{"question_id":"purchase_intent","value":""}', "空字符串 value"),
    ('{"question_id":"purchase_intent","value":"   "}', "纯空白 value"),
    ('{"question_id":"purchase_intent","value":"UNSURE"}', "大写 UNSURE 不得容错"),
    ('{"question_id":"purchase_intent","value":" unsure"}', "前导空格不得容错"),
    ('{"question_id":"purchase_intent","value":"unsure "}', "尾随空格不得容错"),
    ('{"question_id":"purchase_intent","value":"unsure\\n"}', "内嵌换行不得容错"),
    ('{"question_id":"purchase_intent","value":"Unsure"}', "首字母大写不得容错"),
    # 结构不对
    ('[{"question_id":"purchase_intent","value":"unsure"}]', "纯 JSON 数组"),
    ('好的，我的答案是 {"question_id":"purchase_intent","value":"unsure"}', "散文包裹 JSON"),
    (
        '```\n```json\n{"question_id":"purchase_intent","value":"unsure"}\n```\n```',
        "嵌套两层围栏（只能剥一层）",
    ),
    ('{"value":"unsure"}', "缺 question_id"),
    ('{"question_id":"purchase_intent"}', "缺 value"),
    (
        '{"question_id":"purchase_intent","value":"unsure","confidence":0.9}',
        "额外字段 confidence",
    ),
    ('{"question_id":"purchase_intent","value":"prob', "截断 JSON（模拟 256 token 截断）"),
    ('{"question_id":"purchase_intent","value":"unsure"', "截断 JSON（缺右括号）"),
    ("   \n  ", "空白字符串"),
    ("", "空字符串"),
    (None, "raw_text=None"),
    ("I think probably yes", "自然语言非 JSON"),
]


@pytest.mark.parametrize("raw,desc", _ILLEGAL_OUTPUTS, ids=[d for _, d in _ILLEGAL_OUTPUTS])
def test_qa_illegal_outputs_raise_never_return(raw, desc) -> None:
    """A: 所有无效输出必须抛 InvalidOutputError，**绝不**返回任何（默认）答案。"""
    with pytest.raises(InvalidOutputError) as excinfo:
        parse_and_validate(raw)
    assert excinfo.value.code == "INVALID_OUTPUT"


def test_qa_single_layer_fence_valid_passes() -> None:
    """A: 单层围栏且内容为合法 JSON —— 这是**唯一**允许的容错，应通过。"""
    raw = '```json\n{"question_id":"purchase_intent","value":"definitely_yes"}\n```'
    answer = parse_and_validate(raw)
    assert answer.value == "definitely_yes"
    assert answer.score == 5
    # 无语言标签的单层围栏也应通过
    assert (
        parse_and_validate('```\n{"question_id":"purchase_intent","value":"unsure"}\n```').value
        == "unsure"
    )


def test_qa_fence_with_trailing_text_not_tolerated() -> None:
    """A: 围栏后还有散文 → 不得容错（不允许正则猜答案）。"""
    raw = '```json\n{"question_id":"purchase_intent","value":"unsure"}\n```\n以上是答案'
    with pytest.raises(InvalidOutputError):
        parse_and_validate(raw)


def test_qa_score_derived_from_value_model_score_ignored() -> None:
    """A3: 模型传 score 必须被当作额外字段拒绝；score 只能由 value 派生。"""
    # 模型试图用 score:5 配 value:definitely_not —— 不能让它生效
    with pytest.raises(InvalidOutputError):
        parse_and_validate('{"question_id":"purchase_intent","value":"definitely_not","score":5}')
    with pytest.raises(ValidationError):
        ValidatedAnswer(  # type: ignore[call-arg]
            question_id="purchase_intent", value="definitely_not", score=5
        )
    # 派生关系与主文档 §5.1 完全一致
    for value, score in OPTION_SCORES.items():
        assert ValidatedAnswer(question_id="purchase_intent", value=value).score == score


# ===========================================================================
# G. 契约完整性（T0）
# ===========================================================================


@pytest.mark.parametrize(
    "model,expected",
    [
        (PersonaSnapshot, "PersonaSnapshot"),
        (SurveyInput, "SurveyInput"),
        (RunCreate, "RunCreate"),
        (ModelRequest, "ModelRequest"),
        (ModelResponse, "ModelResponse"),
        (ValidatedAnswer, "ValidatedAnswer"),
    ],
)
def test_qa_six_required_models_exist_verbatim(model, expected) -> None:
    """G1: 六个强制命名模型必须存在且命名逐字正确。"""
    assert model.__name__ == expected


def test_qa_five_values_exact_and_no_reorder_or_dup() -> None:
    """G2: 五档 value 必须**恰好**是固定顺序集合：乱序/重复/增减均被拒。"""
    assert OPTION_VALUES == (
        "definitely_not",
        "probably_not",
        "unsure",
        "probably_yes",
        "definitely_yes",
    )
    # 乱序
    reordered = default_options()[::-1]
    with pytest.raises(ValidationError):
        QuestionInput(options=reordered)
    # 重复（P2-a 起 QuestionOption.label 为 Literal，占位 label 改用规范中文展示）
    dup = default_options()
    dup[1] = QuestionOption(value="definitely_not", label="肯定不会购买", score=1)
    with pytest.raises(ValidationError):
        QuestionInput(options=dup)
    # 少一个
    with pytest.raises(ValidationError):
        QuestionInput(options=default_options()[:4])


def test_qa_option_triple_must_consistently_match() -> None:
    """P2-a 回归：value↔label↔score 三元组错配 + 自由文本 label 必须被拒。

    意图：锁定 ``QuestionOption`` 的纵深防御（避免错配 label/score 静默污染报表）。
    """
    # label 与 value 不符
    with pytest.raises(ValidationError):
        QuestionOption(value="definitely_not", label="肯定会购买", score=1)
    # score 与 value 不符
    with pytest.raises(ValidationError):
        QuestionOption(value="definitely_not", label="肯定不会购买", score=5)
    # 自由文本 label 被拒（原 P2-a 缺陷点）
    with pytest.raises(ValidationError):
        QuestionOption(value="definitely_not", label="X", score=1)
    # 规范三元组通过
    ok = QuestionOption(value="probably_yes", label="可能会购买", score=4)
    assert ok.score == 4


def test_qa_question_id_and_type_fixed() -> None:
    """G: question_id / type 由 Literal 类型层固定，不可改。"""
    with pytest.raises(ValidationError):
        QuestionInput(question_id="other")  # type: ignore[arg-type]
    with pytest.raises(ValidationError):
        QuestionInput(type="likert")  # type: ignore[arg-type]
    assert QuestionInput().question_id == PURCHASE_INTENT_QUESTION_ID


def test_qa_single_question_structurally_enforced() -> None:
    """G2: 只允许 1 道题 —— 列表式/复数式题目必须被拒（extra=forbid）。"""
    product = ProductInput(
        name="p", description="d", price=Decimal("1"), price_unit="元"
    )
    with pytest.raises(ValidationError):
        SurveyInput(title="t", product=product, question=[QuestionInput()])  # type: ignore[arg-type]
    with pytest.raises(ValidationError):
        SurveyInput(  # type: ignore[call-arg]
            title="t", product=product, question=QuestionInput(), questions=[QuestionInput()]
        )


def test_qa_reject_missing_price_persona_and_profile() -> None:
    """G4: 缺价格 / 缺画像 ID / 缺画像文本必须被拒。"""
    with pytest.raises(ValidationError):
        ProductInput(name="p", description="d", price_unit="元")  # 缺 price
    with pytest.raises(ValidationError):
        PersonaSnapshot(row_no=1, profile_text="x")  # 缺 persona_id
    with pytest.raises(ValidationError):
        PersonaSnapshot(row_no=1, persona_id="p")  # 缺 profile_text


def test_qa_no_auto_age_filter_in_build_snapshots() -> None:
    """G5: 导入层不得按年龄等字段自动过滤输入行（全量保留）。"""
    rows = [
        HEADER,
        ["p1", 15, "上海", "女", "a"],
        ["p2", "", "北京", "男", "b"],  # 缺年龄 → 不得因此丢行
        ["p3", 99, "广州", "女", "c"],
        ["p4", 0, "深圳", "男", "d"],
    ]
    snaps = build_snapshots(rows, column_mapping=CSV_MAPPING)
    assert [s.persona_id for s in snaps] == ["p1", "p2", "p3", "p4"]
    assert [s.row_no for s in snaps] == [1, 2, 3, 4]
    assert snaps[1].age is None


# ===========================================================================
# C. 数据导入：绝不静默丢行 + 位置化错误
# ===========================================================================


def test_qa_bad_row3_empty_persona_rejects_whole_batch_with_position() -> None:
    """C1: 第 3 行 persona_id 为空 → **整批被拒**（不是跳过坏行），并给出行号+列名。"""
    rows = [HEADER] + [[f"p{i}", 20 + i, "上海", "女", "a"] for i in range(5)]
    rows[3] = ["", 23, "上海", "女", "a"]  # 第 3 数据行 persona_id 为空
    with pytest.raises(InvalidUserTableError) as excinfo:
        parse_user_table(_csv_bytes(rows[0], rows[1:]), column_mapping=CSV_MAPPING)
    error = excinfo.value
    assert error.code == "INVALID_USER_TABLE"
    assert error.row_no == 3
    assert error.column == "user_id"
    details = error.to_details()
    assert details["row_no"] == 3 and details["column"] == "user_id"


def test_qa_duplicate_persona_id_whole_batch_rejected() -> None:
    """C2: 重复 persona_id → 整批被拒，并给出行号与列名。"""
    rows = [HEADER, ["p1", 20, "上海", "女", "a"], ["p1", 21, "北京", "男", "b"]]
    with pytest.raises(InvalidUserTableError) as excinfo:
        parse_user_table(_csv_bytes(rows[0], rows[1:]), column_mapping=CSV_MAPPING)
    assert excinfo.value.row_no == 2
    assert excinfo.value.column == "user_id"


def test_qa_empty_table_and_empty_profile_rejected() -> None:
    """C2: 空表 / 空画像 → 被拒（不得产出 0 行冲抵）。"""
    with pytest.raises(InvalidUserTableError):
        parse_user_table(_csv_bytes(HEADER, []), column_mapping=CSV_MAPPING)
    with pytest.raises(InvalidUserTableError) as excinfo:
        parse_user_table(
            _csv_bytes(HEADER, [["p1", "", "", "", ""]]), column_mapping=CSV_MAPPING
        )
    assert excinfo.value.row_no == 1


def test_qa_over_20000_rows_rejected() -> None:
    """C2: 20,001 行 → 被拒（上限 20,000）。"""
    rows = [[f"p{i}", 20, "上海", "女", "a"] for i in range(20_001)]
    with pytest.raises(InvalidUserTableError) as excinfo:
        parse_user_table(_csv_bytes(HEADER, rows), column_mapping=CSV_MAPPING)
    assert excinfo.value.code == "INVALID_USER_TABLE"


def test_qa_csv_bom_and_no_bom_both_parsed() -> None:
    """C4: 带 BOM 与不带 BOM 的 UTF-8 CSV 都要能读。"""
    rows = [["p1", 20, "上海", "女", "a"], ["p2", 21, "北京", "男", "b"]]
    with_bom = parse_user_table(
        _csv_bytes(HEADER, rows, bom=True), column_mapping=CSV_MAPPING
    )
    without_bom = parse_user_table(
        _csv_bytes(HEADER, rows, bom=False), column_mapping=CSV_MAPPING
    )
    assert [s.persona_id for s in with_bom.snapshots] == ["p1", "p2"]
    assert [s.persona_id for s in without_bom.snapshots] == ["p1", "p2"]


def test_qa_row_order_preserved_not_sorted() -> None:
    """C6: 行顺序按输入顺序保留（不是字典序/随机序）。"""
    ids = ["p_030", "p_001", "p_999", "p_abc"]
    rows = [[pid, 30, "上海", "女", "n"] for pid in ids]
    parsed = parse_user_table(_csv_bytes(HEADER, rows), column_mapping=CSV_MAPPING)
    assert [s.persona_id for s in parsed.snapshots] == ids
    assert [s.row_no for s in parsed.snapshots] == [1, 2, 3, 4]


# ===========================================================================
# C5. XLSX：工作表选择 + 公式不执行（映射字段含公式须要求转静态值）
# ===========================================================================


def _xlsx_bytes(sheets: dict[str, tuple[list, list[list]]]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, (header, data_rows) in sheets.items():
        ws = workbook.create_sheet(title=name)
        ws.append(header)
        for row in data_rows:
            ws.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_qa_xlsx_sheet_selection() -> None:
    """C5: 能选中指定工作表；默认第一张；不存在的工作表报错。"""
    payload = _xlsx_bytes(
        {
            "first": (HEADER, [["p_first", 20, "上海", "女", "a"]]),
            "second": (HEADER, [["p_second", 21, "北京", "男", "b"]]),
        }
    )
    assert parse_user_table(payload, column_mapping=CSV_MAPPING).sheet_name == "first"
    second = parse_user_table(payload, sheet_name="second", column_mapping=CSV_MAPPING)
    assert [s.persona_id for s in second.snapshots] == ["p_second"]
    with pytest.raises(InvalidUserTableError):
        parse_user_table(payload, sheet_name="nope", column_mapping=CSV_MAPPING)


def test_qa_xlsx_formula_in_mapped_field_rejected() -> None:
    """C5: 映射字段里的公式单元格必须被拒（不执行公式，要求转静态值后重传）。"""
    from openpyxl import Workbook

    workbook = Workbook()
    ws = workbook.active
    ws.title = "personas"
    ws.append(HEADER)
    ws.append(["p1", "=1+1", "上海", "女", "a"])  # age 是公式
    buffer = io.BytesIO()
    workbook.save(buffer)

    with pytest.raises(InvalidUserTableError) as excinfo:
        parse_user_table(buffer.getvalue(), column_mapping=CSV_MAPPING)
    assert excinfo.value.code == "INVALID_USER_TABLE"
    assert "formula" in excinfo.value.message.lower()
    # 位置信息可用（列名/行号）
    assert excinfo.value.column == "age" or excinfo.value.row_no == 1


@pytest.mark.parametrize(
    "label,data,filename",
    [
        ("corrupt-xlsx-by-magic", b"PK\x03\x04not-a-real-zip", None),
        ("xlsx-filename-garbage", b"garbage-bytes-here", "users.xlsx"),
        ("binary-csv", b"\xff\xfe\x00\x01\x02\x03bad", None),
    ],
    ids=["corrupt-xlsx", "xlsx-filename-garbage", "binary-csv"],
)
def test_qa_malformed_file_maps_to_invalid_user_table(label, data, filename) -> None:
    """P1 回归：损坏 XLSX / 非 UTF-8 CSV 必须映射为 INVALID_USER_TABLE（不得抛裸异常→500）。"""
    with pytest.raises(InvalidUserTableError) as excinfo:
        parse_user_table(data, filename=filename, column_mapping=CSV_MAPPING)
    assert excinfo.value.code == "INVALID_USER_TABLE"


# ===========================================================================
# C3. **真·20,000 行进库** —— 并用 SQL 直接复核（工程师只测了纯解析，未测入库）
# ===========================================================================


@pytest.mark.asyncio
async def test_qa_20000_real_db_import_verified_by_sql(db_session) -> None:
    """C3: 20,000 行必须**真进库**：row_count、快照条数、row_no 连续无缺号（SQL 复核）。

    意图：工程师的 20,000 用例只调用 ``parse_user_table``（纯解析），
    从未验证 DB 落库路径。本用例直查 ``imports`` 表，并用
    ``jsonb_array_length`` 与 ``jsonb_array_elements`` 复核 row_no 的
    min/max/count/distinct，堵住「fixture 其实没真导入」的可能。
    """
    from tests.fixtures.generate_personas import generate_persona_rows

    rows = generate_persona_rows(20_000)
    data = _csv_bytes(HEADER, [[r[c] for c in HEADER] for r in rows])

    record, parsed = await create_import(
        db_session, data=data, filename="p20k.csv", column_mapping=CSV_MAPPING
    )
    await db_session.commit()

    # (1) ORM 侧
    assert record.row_count == 20_000
    assert len(parsed.snapshots) == 20_000

    # (2) 直接查库：row_count 列 + 快照 JSONB 数组长度
    db_row_count = await db_session.scalar(
        text("SELECT row_count FROM imports WHERE id = :i"), {"i": record.id}
    )
    db_snap_len = await db_session.scalar(
        text("SELECT jsonb_array_length(snapshots_json) FROM imports WHERE id = :i"),
        {"i": record.id},
    )
    assert db_row_count == 20_000
    assert db_snap_len == 20_000

    # (3) 用 SQL 展开快照，核对 row_no 的 min/max/count/distinct（连续无缺号）
    stats = (
        await db_session.execute(
            text(
                """
                SELECT min((e->>'row_no')::int) AS mn,
                       max((e->>'row_no')::int) AS mx,
                       count(*)                AS cnt,
                       count(DISTINCT (e->>'row_no')::int) AS dst,
                       count(DISTINCT (e->>'persona_id')) AS pid_cnt
                FROM imports, jsonb_array_elements(snapshots_json) AS e
                WHERE id = :i
                """
            ),
            {"i": record.id},
        )
    ).one()
    assert stats.mn == 1 and stats.mx == 20_000
    assert stats.cnt == 20_000 and stats.dst == 20_000
    assert stats.pid_cnt == 20_000

    # (4) load_snapshots 读回后顺序与 row_no 仍连续
    from app.personas.source import load_snapshots

    loaded = await load_snapshots(db_session, record.id)
    assert [s.row_no for s in loaded] == list(range(1, 20_001))
    assert loaded[0].persona_id == "p_000001"


# ===========================================================================
# D. 测试库隔离（运行时断言）
# ===========================================================================


@pytest.mark.asyncio
async def test_qa_running_against_isolated_test_db(db_session) -> None:
    """D: 运行时必须连在以 survey_test 开头的库上（fail fast 由 conftest 保证，这里再兜一层）。

    裁定 §7.6.4：放宽为前缀匹配，以支持 ``survey_test_qa`` 等有意隔离的库。
    """
    name = await db_session.scalar(text("SELECT current_database()"))
    assert str(name).startswith("survey_test")
