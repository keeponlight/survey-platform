# SPDX-License-Identifier: GPL-3.0-or-later
"""T1 用户表导入测试（主文档 §4 / task-list.md T1）。

验收命令：``uv run pytest tests/test_personas.py tests/test_surveys.py -q``
"""

from __future__ import annotations

import csv
import io
import sys

import pytest
from openpyxl import Workbook

from app.contracts import ERROR_INVALID_USER_TABLE
from app.personas.source import (
    InvalidUserTableError,
    build_preview,
    create_import,
    load_snapshots,
    parse_user_table,
)
from tests.fixtures.generate_personas import (
    DEFAULT_COLUMN_MAPPING,
    generate_persona_rows,
    write_personas_csv,
)

CSV_MAPPING = DEFAULT_COLUMN_MAPPING


def _csv_bytes(header: list[str], data_rows: list[list[object]], *, bom: bool = False) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(header)
    for row in data_rows:
        writer.writerow(row)
    encoding = "utf-8-sig" if bom else "utf-8"
    return buffer.getvalue().encode(encoding)


def _xlsx_bytes(sheets: dict[str, tuple[list[str], list[list[object]]]]) -> bytes:
    workbook = Workbook()
    # 删除默认空表，按给定顺序创建。
    default = workbook.active
    workbook.remove(default)
    for name, (header, data_rows) in sheets.items():
        worksheet = workbook.create_sheet(title=name)
        worksheet.append(header)
        for row in data_rows:
            worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# 版本钉定（比只验 (3,12) 更严）
# ---------------------------------------------------------------------------


def test_interpreter_matches_pinned_patch() -> None:
    """解释器必须精确为 (3, 12, 13)（backend/.python-version 钉版生效）。"""
    assert sys.version_info[:3] == (3, 12, 13), (
        f"expected Python 3.12.13, got {sys.version_info[:3]} "
        "(check backend/.python-version and that uv resolved it)"
    )


# ---------------------------------------------------------------------------
# 行序 / 无静默丢行
# ---------------------------------------------------------------------------


def test_row_order_preserved() -> None:
    """行顺序与 row_no 必须原样保留（1 起），persona_id 不重排。"""
    header = ["user_id", "age", "city", "gender", "note"]
    ids = ["p_030", "p_001", "p_999", "p_abc"]
    rows = [[pid, 30, "上海", "女", "备注"] for pid in ids]
    parsed = parse_user_table(_csv_bytes(header, rows), column_mapping=CSV_MAPPING)

    assert [snapshot.persona_id for snapshot in parsed.snapshots] == ids
    assert [snapshot.row_no for snapshot in parsed.snapshots] == [1, 2, 3, 4]


def test_no_silent_row_drop() -> None:
    """干净表：行数不丢；坏行：必须报错而不是被静默跳过。"""
    header = ["user_id", "age", "city", "gender", "note"]
    good_rows = [[f"p_{index}", 25, "北京", "男", "备注"] for index in range(5)]
    parsed = parse_user_table(_csv_bytes(header, good_rows), column_mapping=CSV_MAPPING)
    assert len(parsed.snapshots) == 5

    # 第 3 行画像为空 → 必须报错（而不是得到 4 行结果）。
    bad_rows = [[f"p_{index}", 25, "", "", ""] for index in range(5)]
    bad_rows[2] = ["p_2", "", "", "", ""]  # 空画像
    with pytest.raises(InvalidUserTableError) as excinfo:
        parse_user_table(_csv_bytes(header, bad_rows), column_mapping=CSV_MAPPING)
    assert excinfo.value.code == ERROR_INVALID_USER_TABLE
    assert excinfo.value.row_no == 3


# ---------------------------------------------------------------------------
# 位置化错误
# ---------------------------------------------------------------------------


def test_duplicate_id_rejected_with_position() -> None:
    """重复 persona_id 必须被拒，并给出行号 + 列名。"""
    header = ["user_id", "age", "city", "gender", "note"]
    rows = [
        ["p_1", 20, "上海", "女", "a"],
        ["p_2", 21, "上海", "女", "b"],
        ["p_1", 22, "上海", "女", "c"],  # 与第 1 行重复
    ]
    with pytest.raises(InvalidUserTableError) as excinfo:
        parse_user_table(_csv_bytes(header, rows), column_mapping=CSV_MAPPING)
    error = excinfo.value
    assert error.code == ERROR_INVALID_USER_TABLE
    assert error.row_no == 3
    assert error.column == "user_id"
    details = error.to_details()
    assert details["row_no"] == 3
    assert details["duplicate_of_row"] == 1


def test_empty_profile_rejected() -> None:
    """空画像必须被拒，并给出行号 + 画像列名。"""
    header = ["user_id", "age", "city", "gender", "note"]
    rows = [["p_1", 20, "上海", "女", "a"], ["p_2", "", "", "", ""]]
    with pytest.raises(InvalidUserTableError) as excinfo:
        parse_user_table(_csv_bytes(header, rows), column_mapping=CSV_MAPPING)
    error = excinfo.value
    assert error.row_no == 2
    assert "age" in (error.column or "")


def test_empty_table_rejected() -> None:
    header = ["user_id", "age", "city", "gender", "note"]
    with pytest.raises(InvalidUserTableError) as excinfo:
        parse_user_table(_csv_bytes(header, []), column_mapping=CSV_MAPPING)
    assert excinfo.value.code == ERROR_INVALID_USER_TABLE


def test_missing_mapped_column_reports_column() -> None:
    header = ["user_id", "age"]  # 缺少 city/gender/note
    rows = [["p_1", 20]]
    with pytest.raises(InvalidUserTableError) as excinfo:
        parse_user_table(_csv_bytes(header, rows), column_mapping=CSV_MAPPING)
    assert excinfo.value.column in {"city", "gender", "note"}


def test_too_many_rows_rejected() -> None:
    header = ["user_id", "age", "city", "gender", "note"]
    rows = [[f"p_{index}", 20, "上海", "女", "a"] for index in range(20_001)]
    with pytest.raises(InvalidUserTableError) as excinfo:
        parse_user_table(_csv_bytes(header, rows), column_mapping=CSV_MAPPING)
    assert excinfo.value.code == ERROR_INVALID_USER_TABLE


# ---------------------------------------------------------------------------
# CSV BOM / XLSX 工作表
# ---------------------------------------------------------------------------


def test_csv_with_bom_parsed() -> None:
    """UTF-8 带 BOM 的 CSV 必须能解析（utf-8-sig），且首列名不被 BOM 污染。"""
    header = ["user_id", "age", "city", "gender", "note"]
    rows = [["p_1", 20, "上海", "女", "a"], ["p_2", 21, "北京", "男", "b"]]
    parsed = parse_user_table(_csv_bytes(header, rows, bom=True), column_mapping=CSV_MAPPING)
    assert [s.persona_id for s in parsed.snapshots] == ["p_1", "p_2"]


def test_xlsx_sheet_selection() -> None:
    """XLSX 必须支持用户指定工作表（默认第一张）。"""
    header = ["user_id", "age", "city", "gender", "note"]
    payload = _xlsx_bytes(
        {
            "first": (header, [["p_first", 20, "上海", "女", "a"]]),
            "second": (header, [["p_second", 21, "北京", "男", "b"]]),
        }
    )

    # 默认第一张
    default_parsed = parse_user_table(payload, column_mapping=CSV_MAPPING)
    assert default_parsed.sheet_name == "first"
    assert [s.persona_id for s in default_parsed.snapshots] == ["p_first"]

    # 显式指定第二张
    second_parsed = parse_user_table(payload, sheet_name="second", column_mapping=CSV_MAPPING)
    assert second_parsed.sheet_name == "second"
    assert [s.persona_id for s in second_parsed.snapshots] == ["p_second"]

    # 指定不存在的工作表 → 报错并列出可用工作表
    with pytest.raises(InvalidUserTableError) as excinfo:
        parse_user_table(payload, sheet_name="missing", column_mapping=CSV_MAPPING)
    assert excinfo.value.code == ERROR_INVALID_USER_TABLE


def test_xlsx_formula_cell_rejected() -> None:
    """XLSX 不执行公式：映射字段里的公式单元格必须被拒（要求转静态值）。"""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "personas"
    worksheet.append(["user_id", "age", "city", "gender", "note"])
    worksheet.append(["p_1", "=1+1", "上海", "女", "a"])  # age 是公式
    buffer = io.BytesIO()
    workbook.save(buffer)

    with pytest.raises(InvalidUserTableError) as excinfo:
        parse_user_table(buffer.getvalue(), column_mapping=CSV_MAPPING)
    assert excinfo.value.code == ERROR_INVALID_USER_TABLE
    assert "formula" in excinfo.value.message.lower()


# ---------------------------------------------------------------------------
# 20,000 行开发 fixture
# ---------------------------------------------------------------------------


def test_20000_fixture_import_rowcount(tmp_path) -> None:
    """20,000 行 fixture 必须完整导入，行数与行号正确、无静默丢行。"""
    csv_path = write_personas_csv(tmp_path / "personas_20000.csv", 20_000)
    data = csv_path.read_bytes()
    parsed = parse_user_table(
        data, filename=csv_path.name, column_mapping=CSV_MAPPING
    )

    assert len(parsed.snapshots) == 20_000
    assert parsed.snapshots[0].row_no == 1
    assert parsed.snapshots[0].persona_id == "p_000001"
    assert parsed.snapshots[-1].row_no == 20_000
    assert parsed.snapshots[-1].persona_id == "p_020000"
    # 行号连续且唯一
    assert [s.row_no for s in parsed.snapshots] == list(range(1, 20_001))
    # persona_id 唯一
    assert len({s.persona_id for s in parsed.snapshots}) == 20_000


def test_fixture_generator_is_deterministic() -> None:
    assert generate_persona_rows(5) == generate_persona_rows(5)
    assert generate_persona_rows(3, seed=1) != generate_persona_rows(3, seed=2)


# ---------------------------------------------------------------------------
# DB 落库 + 预览 + 测试库隔离
# ---------------------------------------------------------------------------


async def test_import_record_roundtrip_and_preview(db_session) -> None:
    """create_import 落库 → load_snapshots 读回；预览 3 行；source_version 冻结。"""
    header = ["user_id", "age", "city", "gender", "note"]
    rows = [[f"p_{index}", 20 + index, "上海", "女", "a"] for index in range(5)]
    data = _csv_bytes(header, rows)

    record, parsed = await create_import(
        db_session, data=data, filename="t.csv", column_mapping=CSV_MAPPING
    )
    await db_session.commit()

    assert record.row_count == 5
    assert record.file_hash
    preview = build_preview(record, parsed)
    assert preview.row_count == 5
    assert len(preview.preview) == 3
    assert preview.preview[0].persona_id == "p_0"

    loaded = await load_snapshots(db_session, record.id)
    assert [s.persona_id for s in loaded] == [f"p_{index}" for index in range(5)]
    assert loaded[0].source_version == str(record.id)


async def test_test_db_is_separated(db_session) -> None:
    """测试库隔离断言：当前数据库必须以 survey_test 开头（防止误连开发/生产库）。

    裁定 §7.6.4：放宽为前缀匹配，允许 ``survey_test_qa`` / ``survey_test_t7`` 等
    有意隔离的库；dev 库 ``survey`` 不匹配前缀，仍被拒绝。
    """
    from sqlalchemy import text

    result = await db_session.execute(text("SELECT current_database()"))
    assert str(result.scalar_one()).startswith("survey_test")


# ---------------------------------------------------------------------------
# P1：文件格式错误必须映射为 INVALID_USER_TABLE（不得泄漏裸异常）
# ---------------------------------------------------------------------------


def test_binary_csv_is_invalid_user_table() -> None:
    """二进制内容当文本 CSV 传入 → InvalidUserTableError（不是 UnicodeDecodeError）。"""
    binary = b"\x80\x81\x82\xff\xfe\x00\xc3\x28"  # 非法 UTF-8 字节序列
    with pytest.raises(InvalidUserTableError) as excinfo:
        parse_user_table(binary, column_mapping=CSV_MAPPING)
    assert excinfo.value.code == ERROR_INVALID_USER_TABLE
    # 明确不是裸解码异常
    assert not isinstance(excinfo.value, UnicodeDecodeError)


def test_garbage_xlsx_is_invalid_user_table() -> None:
    """随机字节伪装成 .xlsx → InvalidUserTableError（不是 BadZipFile）。"""
    import zipfile

    garbage = b"this is definitely not a zip workbook"
    with pytest.raises(InvalidUserTableError) as excinfo:
        parse_user_table(garbage, filename="broken.xlsx", column_mapping=CSV_MAPPING)
    assert excinfo.value.code == ERROR_INVALID_USER_TABLE
    assert not isinstance(excinfo.value, zipfile.BadZipFile)


def test_truncated_xlsx_is_invalid_user_table() -> None:
    """截断的真实 XLSX → InvalidUserTableError（不是 BadZipFile）。"""
    import zipfile

    header = ["user_id", "age", "city", "gender", "note"]
    full = _xlsx_bytes({"personas": (header, [["p_1", 20, "上海", "女", "a"]])})
    truncated = full[: len(full) // 2]
    with pytest.raises(InvalidUserTableError) as excinfo:
        parse_user_table(truncated, filename="trunc.xlsx", column_mapping=CSV_MAPPING)
    assert excinfo.value.code == ERROR_INVALID_USER_TABLE
    assert not isinstance(excinfo.value, zipfile.BadZipFile)


def test_format_error_not_swallowed_as_default() -> None:
    """损坏文件**绝不返回任何行**：返回路径必须是异常（不「猜」数据）。"""
    # 二进制 CSV：绝不返回被「勉强解码」出来的行
    result = None
    try:
        result = parse_user_table(b"\xff\xfe\x00\x81", column_mapping=CSV_MAPPING)
    except InvalidUserTableError:
        result = None
    assert result is None

    # 损坏 XLSX：同样必须抛错，绝不返回批次数据
    with pytest.raises(InvalidUserTableError):
        parse_user_table(b"PK\x03\x04nope", filename="x.xlsx", column_mapping=CSV_MAPPING)


def test_empty_file_still_invalid_user_table_at_public_layer() -> None:
    """空文件在公开层仍为 InvalidUserTableError（P1 未破坏既有语义）。"""
    with pytest.raises(InvalidUserTableError) as excinfo:
        parse_user_table(b"", column_mapping=CSV_MAPPING)
    assert excinfo.value.code == ERROR_INVALID_USER_TABLE

    # 低层私有函数对空输入返回空列表（分层合理，未被改动）。
    from app.personas.source import _read_csv_rows

    assert _read_csv_rows(b"") == (None, [])

