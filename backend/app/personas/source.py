# SPDX-License-Identifier: GPL-3.0-or-later
"""用户表导入、列映射与逐行快照（主文档 §4）。

硬约束：
- 支持 CSV（UTF-8，可带 BOM，用 ``utf-8-sig``）与 XLSX（用户指定工作表，默认第一张）。
- **保留** ``row_no`` 与 ``persona_id``；**不按年龄等字段自动过滤输入行**。
- 重复 ID / 空画像 / 空表 / 超 20,000 行 → ``INVALID_USER_TABLE``，并给出
  **工作表名 / 行号 / 列名**；**绝不静默丢行**。
- XLSX **不执行公式**：映射字段中的公式单元格要求转为静态值后重传。
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import zipfile
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from typing import Any
from uuid import UUID, uuid4

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.contracts import ERROR_INVALID_USER_TABLE, PersonaSnapshot
from app.models import ImportRecord

#: 上传大小上限（主文档 §4：20 MB）。
MAX_FILE_BYTES = 20 * 1024 * 1024

#: 最大数据行数（主文档 §4：20,000）。
MAX_ROWS = 20_000

#: 列映射中必须提供的键。
REQUIRED_MAPPING_KEYS: tuple[str, ...] = ("persona_id", "profile_text_columns")

#: 可选的画像附加字段（映射到 PersonaSnapshot 的同名可选字段）。
OPTIONAL_SINGLE_COLUMNS: tuple[str, ...] = ("age", "city_tier")


class InvalidUserTableError(Exception):
    """用户表不合法（错误码 ``INVALID_USER_TABLE``）。

    ``details`` 至少包含出错位置（工作表/行号/列名）之一，供 API 原样返回给用户修复。
    """

    code = ERROR_INVALID_USER_TABLE

    def __init__(
        self,
        message: str,
        *,
        sheet: str | None = None,
        row_no: int | None = None,
        column: str | None = None,
        extra: Mapping[str, Any] | None = None,
    ) -> None:
        super().__init__(message)
        self.message = message
        self.sheet = sheet
        self.row_no = row_no
        self.column = column
        self.extra = dict(extra or {})

    def to_details(self) -> dict[str, Any]:
        details: dict[str, Any] = {"message": self.message}
        if self.sheet is not None:
            details["sheet"] = self.sheet
        if self.row_no is not None:
            details["row_no"] = self.row_no
        if self.column is not None:
            details["column"] = self.column
        details.update(self.extra)
        return details


class ImportPreview(BaseModel):
    """导入预览（主文档 §4：import_id、row_count、3 行预览、行级错误）。"""

    model_config = ConfigDict(extra="forbid")

    import_id: UUID
    row_count: int
    sheet_name: str | None = None
    source_version: str
    preview: list[PersonaSnapshot] = Field(default_factory=list)
    errors: list[dict[str, Any]] = Field(default_factory=list)


@dataclass(frozen=True)
class ParsedTable:
    """解析结果（不含 DB 交互）。"""

    snapshots: list[PersonaSnapshot]
    sheet_name: str | None
    headers: list[str]


# ---------------------------------------------------------------------------
# 底层解析
# ---------------------------------------------------------------------------


def _read_csv_rows(data: bytes) -> tuple[str | None, list[list[str]]]:
    """读 CSV（UTF-8，兼容 BOM）。返回 (sheet_name=None, 行列表)。

    非 UTF-8 文本 → :class:`InvalidUserTableError`（**不**猜测编码、**不**用
    ``errors="ignore"/"replace"`` 硬解，避免「猜」出一批不存在的数据）。
    空字节输入仍返回空列表 —— 「空表」由公开层 :func:`parse_user_table` 判定。
    """
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise InvalidUserTableError(
            "文件不是有效的 UTF-8 CSV/文本（file is not valid UTF-8 CSV/text）",
            extra={"reason": type(exc).__name__},
        ) from exc
    reader = csv.reader(io.StringIO(text))
    rows = [list(row) for row in reader]
    return None, rows


def _read_xlsx_rows(data: bytes, sheet_name: str | None) -> tuple[str | None, list[list[Any]]]:
    """读 XLSX 指定工作表（默认第一张）。返回 (实际工作表名, 行列表)。

    非 XLSX / 损坏的 zip → :class:`InvalidUserTableError`（不泄漏裸异常，
    避免 T4 把「文件格式错误」错映射成 500）。

    openpyxl 默认 ``data_only=False``，公式单元格返回公式字符串（如 ``=A1+B1``），
    本函数**不执行公式**；公式检测在 :func:`_cell_to_text` 内完成。
    """
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError) as exc:
        raise InvalidUserTableError(
            "文件不是有效的 XLSX 工作簿（file is not a valid XLSX workbook）",
            sheet=sheet_name,
            extra={"reason": type(exc).__name__},
        ) from exc
    try:
        if sheet_name is not None:
            if sheet_name not in workbook.sheetnames:
                raise InvalidUserTableError(
                    f"worksheet not found: {sheet_name!r}",
                    sheet=sheet_name,
                    extra={"available_sheets": list(workbook.sheetnames)},
                )
            worksheet = workbook[sheet_name]
            resolved_name = sheet_name
        else:
            resolved_name = workbook.sheetnames[0] if workbook.sheetnames else None
            if resolved_name is None:
                raise InvalidUserTableError("workbook has no worksheet")
            worksheet = workbook[resolved_name]
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        return resolved_name, rows
    finally:
        workbook.close()


def _cell_to_text(
    value: Any,
    *,
    sheet: str | None,
    row_no: int,
    column: str,
    is_xlsx: bool,
) -> str:
    """把单元格值转为文本；XLSX 公式单元格直接拒绝（不执行公式）。"""
    if value is None:
        return ""
    if is_xlsx and isinstance(value, str) and value.startswith("="):
        raise InvalidUserTableError(
            "formula cell detected; convert to a static value and re-upload",
            sheet=sheet,
            row_no=row_no,
            column=column,
        )
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def _row_cell(
    raw_row: Sequence[Any],
    index: int,
    column: str,
    *,
    sheet: str | None,
    row_no: int,
    is_xlsx: bool,
) -> str:
    """读取某一行某一列的文本（越界返回空串）。"""
    raw_value = raw_row[index] if index < len(raw_row) else None
    return _cell_to_text(
        raw_value, sheet=sheet, row_no=row_no, column=column, is_xlsx=is_xlsx
    )


# ---------------------------------------------------------------------------
# 映射与快照构造
# ---------------------------------------------------------------------------


def _normalise_mapping(column_mapping: Mapping[str, Any]) -> dict[str, Any]:
    if not isinstance(column_mapping, Mapping):
        raise InvalidUserTableError("column_mapping must be an object")
    missing = [key for key in REQUIRED_MAPPING_KEYS if not column_mapping.get(key)]
    if missing:
        raise InvalidUserTableError(
            f"column_mapping missing required keys: {missing}",
            extra={"missing_keys": missing},
        )
    mapping: dict[str, Any] = {
        "persona_id": str(column_mapping["persona_id"]),
        "profile_text_columns": list(column_mapping["profile_text_columns"]),
    }
    for key in OPTIONAL_SINGLE_COLUMNS:
        value = column_mapping.get(key)
        if value:
            mapping[key] = str(value)
    return mapping


def _resolve_index(headers: Sequence[str], column_name: str, *, role: str) -> int:
    try:
        return list(headers).index(column_name)
    except ValueError as exc:
        raise InvalidUserTableError(
            f"column {column_name!r} for role {role!r} not found in table header",
            column=column_name,
            extra={"available_columns": list(headers)},
        ) from exc


def build_snapshots(
    rows: Sequence[Sequence[Any]],
    *,
    column_mapping: Mapping[str, Any],
    sheet_name: str | None = None,
    source_version: str | None = None,
    is_xlsx: bool = False,
) -> list[PersonaSnapshot]:
    """把原始表格行按列映射构造为 ``PersonaSnapshot`` 列表（保留 row_no / persona_id）。

    第一行为表头。数据行 ``row_no`` 从 1 开始（与主文档 §4 示例一致）。
    任何不合法行都会抛 :class:`InvalidUserTableError`，**绝不静默跳过**。
    """
    mapping = _normalise_mapping(column_mapping)

    if len(rows) == 0:
        raise InvalidUserTableError("table is empty (no header row)", sheet=sheet_name)

    header_row = list(rows[0])
    headers = [str(cell).strip() if cell is not None else "" for cell in header_row]
    if not any(headers):
        raise InvalidUserTableError("table header is empty", sheet=sheet_name)

    id_index = _resolve_index(headers, mapping["persona_id"], role="persona_id")
    profile_columns: list[str] = [str(c) for c in mapping["profile_text_columns"]]
    profile_indexes = {
        column: _resolve_index(headers, column, role="profile_text_columns")
        for column in profile_columns
    }
    optional_indexes = {
        key: _resolve_index(headers, mapping[key], role=key)
        for key in OPTIONAL_SINGLE_COLUMNS
        if key in mapping
    }

    data_rows = rows[1:]
    if len(data_rows) == 0:
        raise InvalidUserTableError("table has no data rows", sheet=sheet_name)
    if len(data_rows) > MAX_ROWS:
        raise InvalidUserTableError(
            f"table has {len(data_rows)} data rows, exceeds maximum {MAX_ROWS}",
            sheet=sheet_name,
            extra={"row_count": len(data_rows), "max_rows": MAX_ROWS},
        )

    snapshots: list[PersonaSnapshot] = []
    seen_ids: dict[str, int] = {}

    for offset, raw_row in enumerate(data_rows):
        row_no = offset + 1

        persona_id = _row_cell(
            raw_row,
            id_index,
            mapping["persona_id"],
            sheet=sheet_name,
            row_no=row_no,
            is_xlsx=is_xlsx,
        )
        if not persona_id:
            raise InvalidUserTableError(
                "empty persona id",
                sheet=sheet_name,
                row_no=row_no,
                column=mapping["persona_id"],
            )
        if persona_id in seen_ids:
            raise InvalidUserTableError(
                f"duplicate persona id {persona_id!r} (first seen at row {seen_ids[persona_id]})",
                sheet=sheet_name,
                row_no=row_no,
                column=mapping["persona_id"],
                extra={"duplicate_of_row": seen_ids[persona_id]},
            )
        seen_ids[persona_id] = row_no

        profile_parts: list[str] = []
        for column in profile_columns:
            value = _row_cell(
                raw_row,
                profile_indexes[column],
                column,
                sheet=sheet_name,
                row_no=row_no,
                is_xlsx=is_xlsx,
            )
            if value:
                profile_parts.append(f"{column}: {value}")
        profile_text = "; ".join(profile_parts)
        if not profile_text:
            raise InvalidUserTableError(
                "empty profile text",
                sheet=sheet_name,
                row_no=row_no,
                column=", ".join(profile_columns),
            )

        age: int | None = None
        if "age" in optional_indexes:
            raw_age = _row_cell(
                raw_row,
                optional_indexes["age"],
                mapping["age"],
                sheet=sheet_name,
                row_no=row_no,
                is_xlsx=is_xlsx,
            )
            if raw_age:
                try:
                    age = int(float(raw_age))
                except ValueError as exc:
                    raise InvalidUserTableError(
                        f"invalid age value {raw_age!r}",
                        sheet=sheet_name,
                        row_no=row_no,
                        column=mapping["age"],
                    ) from exc

        city_tier: str | None = None
        if "city_tier" in optional_indexes:
            raw_city = _row_cell(
                raw_row,
                optional_indexes["city_tier"],
                mapping["city_tier"],
                sheet=sheet_name,
                row_no=row_no,
                is_xlsx=is_xlsx,
            )
            city_tier = raw_city or None

        snapshots.append(
            PersonaSnapshot(
                row_no=row_no,
                persona_id=persona_id,
                profile_text=profile_text,
                age=age,
                city_tier=city_tier,
                source_version=source_version,
            )
        )

    return snapshots


def parse_user_table(
    data: bytes,
    *,
    filename: str | None = None,
    sheet_name: str | None = None,
    column_mapping: Mapping[str, Any],
    source_version: str | None = None,
) -> ParsedTable:
    """解析 CSV/XLSX 字节流并返回快照（无 DB 交互，便于单元测试）。"""
    if len(data) > MAX_FILE_BYTES:
        raise InvalidUserTableError(
            f"file too large: {len(data)} bytes, exceeds maximum {MAX_FILE_BYTES}",
            extra={"size_bytes": len(data), "max_bytes": MAX_FILE_BYTES},
        )

    is_xlsx = False
    if filename and filename.lower().endswith((".xlsx", ".xlsm")):
        is_xlsx = True
    elif data[:2] == b"PK":
        # XLSX 是 zip 容器，魔数 PK。
        is_xlsx = True

    if is_xlsx:
        resolved_sheet, rows = _read_xlsx_rows(data, sheet_name)
        snapshots = build_snapshots(
            rows,
            column_mapping=column_mapping,
            sheet_name=resolved_sheet,
            source_version=source_version,
            is_xlsx=True,
        )
        headers = [str(c).strip() if c is not None else "" for c in (rows[0] if rows else [])]
        return ParsedTable(snapshots=snapshots, sheet_name=resolved_sheet, headers=headers)

    resolved_sheet, rows = _read_csv_rows(data)
    snapshots = build_snapshots(
        rows,
        column_mapping=column_mapping,
        sheet_name=resolved_sheet,
        source_version=source_version,
        is_xlsx=False,
    )
    headers = [str(c).strip() if c is not None else "" for c in (rows[0] if rows else [])]
    return ParsedTable(snapshots=snapshots, sheet_name=resolved_sheet, headers=headers)


# ---------------------------------------------------------------------------
# DB 交互（imports 表）
# ---------------------------------------------------------------------------


def file_hash(data: bytes) -> str:
    """内容 hash（原文件解析后可删，保留 hash）。"""
    return hashlib.sha256(data).hexdigest()


def _snapshots_to_json(snapshots: Iterable[PersonaSnapshot]) -> list[dict[str, Any]]:
    return [snapshot.model_dump(mode="json") for snapshot in snapshots]


async def create_import(
    session: AsyncSession,
    *,
    data: bytes,
    filename: str | None = None,
    sheet_name: str | None = None,
    column_mapping: Mapping[str, Any],
    import_id: UUID | None = None,
    source_version: str | None = None,
) -> tuple[ImportRecord, ParsedTable]:
    """解析用户表并在 ``imports`` 表落库一条记录。

    ``source_version`` 默认取新生成的 import uuid（供 run 冻结引用）。
    """
    new_id = import_id or uuid4()
    effective_source_version = source_version or str(new_id)

    parsed = parse_user_table(
        data,
        filename=filename,
        sheet_name=sheet_name,
        column_mapping=column_mapping,
        source_version=effective_source_version,
    )

    record = ImportRecord(
        id=new_id,
        file_hash=file_hash(data),
        sheet_name=parsed.sheet_name,
        column_mapping_json=dict(column_mapping),
        row_count=len(parsed.snapshots),
        snapshots_json=_snapshots_to_json(parsed.snapshots),
    )
    session.add(record)
    await session.flush()
    return record, parsed


async def load_snapshots(session: AsyncSession, import_id: UUID) -> list[PersonaSnapshot]:
    """按 import_id 读回快照（主文档 §4 接口：``load_snapshots(import_id)``）。"""
    result = await session.execute(
        select(ImportRecord).where(ImportRecord.id == import_id)
    )
    record = result.scalar_one_or_none()
    if record is None:
        raise InvalidUserTableError(
            f"import not found: {import_id}", extra={"import_id": str(import_id)}
        )
    return [PersonaSnapshot.model_validate(item) for item in record.snapshots_json]


def build_preview(
    record: ImportRecord,
    parsed: ParsedTable,
    *,
    preview_size: int = 3,
) -> ImportPreview:
    """构造 3 行预览（主文档 §4）。"""
    return ImportPreview(
        import_id=record.id,
        row_count=record.row_count,
        sheet_name=record.sheet_name,
        source_version=parsed.snapshots[0].source_version or str(record.id)
        if parsed.snapshots
        else str(record.id),
        preview=parsed.snapshots[:preview_size],
    )


def snapshots_to_json(snapshots: Sequence[PersonaSnapshot]) -> str:
    """调试辅助：快照列表 → JSON 文本。"""
    return json.dumps(_snapshots_to_json(snapshots), ensure_ascii=False)


__all__ = [
    "MAX_FILE_BYTES",
    "MAX_ROWS",
    "ImportPreview",
    "InvalidUserTableError",
    "ParsedTable",
    "build_preview",
    "build_snapshots",
    "create_import",
    "file_hash",
    "load_snapshots",
    "parse_user_table",
]
